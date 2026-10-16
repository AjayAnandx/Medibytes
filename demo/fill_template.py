"""Stage 6 demo: fill_template(entities, transcript) -> HTML + DOCX bytes.

Demo-scoped from PIPELINE_ARCHITECTURE_REPORT.md Sec 6.
No FerroTERM/ICD lookup, no 99% gate, no DB. Missing -> ___ + RED badge (no block).
Negated -> excluded from active, shown grey DENIED.
"""
import datetime
import io
import json
import os

_HERE = os.path.dirname(os.path.abspath(__file__))

DEDUP = {"khansi": "cough", "bukhar": "fever"}


def _seg_of(source_sentence, segments):
    for s in segments or []:
        if s.get("text", "").strip() == (source_sentence or "").strip():
            return s.get("id"), f"{s.get('start')}-{s.get('end')}s"
        if source_sentence and source_sentence[:30] in s.get("text", ""):
            return s.get("id"), f"{s.get('start')}-{s.get('end')}s"
    return None, ""


class TemplateRenderError(RuntimeError):
    """The discharge note could not be generated; callers must show this, not a document."""


def fill_template(entities_json, transcript_json, template_id="er_discharge", doc_status=None):
    """Render the discharge note: premium single-A4 slot layout + editable DOCX.

    This is the only supported layout. There is NO silent fallback: a missing
    spec/coords/template file, a slot or render error, or a render that drops
    clinical slot values raises TemplateRenderError and returns no html/docx
    (an empty or partial form must never look like a valid note).
    """
    tpl_dir = os.path.join(_HERE, "templates")
    try:
        with open(os.path.join(tpl_dir, f"{template_id}.json"), encoding="utf-8") as f:
            spec = json.load(f)
    except (OSError, ValueError) as e:
        raise TemplateRenderError(f"discharge note NOT generated: template spec "
                                  f"{template_id}.json unreadable ({e})") from e
    if not os.path.isfile(os.path.join(tpl_dir, f"{template_id}.coords.json")):
        raise TemplateRenderError(f"discharge note NOT generated: layout coords file "
                                  f"{template_id}.coords.json is missing")
    try:
        return _fill_premium(spec, entities_json, transcript_json, template_id, doc_status)
    except TemplateRenderError:
        raise
    except Exception as e:
        raise TemplateRenderError(f"discharge note NOT generated: {template_id} layout failed "
                                  f"({type(e).__name__}: {e})") from e


def _fill_classic(spec, entities_json, transcript_json):
    """OBSOLETE classic layout (templates/er_discharge.html.j2) - not called.

    Kept only for a separate cleanup decision. Unsafe as a fallback: it always
    prints the diagnosis as NIL, shows "None extracted." for allergies (no
    documented-NKA vs not-documented distinction, no conflict flags) and has no
    follow-up, structured vitals or clinician corrections.
    """
    from jinja2 import FileSystemLoader, Environment

    tpl_dir = os.path.join(_HERE, "templates")
    env = Environment(loader=FileSystemLoader(tpl_dir))
    tpl = env.get_template("er_discharge.html.j2")

    ents = entities_json or {}
    segs = (transcript_json or {}).get("segments", [])

    # chief complaint: first non-negated symptom, deduped
    seen, chief = set(), {"text": "NIL — not heard", "missing": True, "color": "RED", "confidence": ""}
    for s in ents.get("symptoms", []):
        if s.get("negated"):
            continue
        key = DEDUP.get(s.get("text", "").lower(), s.get("text", "").lower())
        if key in seen:
            continue
        seen.add(key)
        sid, span = _seg_of(s.get("source_sentence"), segs)
        chief = {"text": f"{key} ({s.get('source_sentence','')[:80]})",
                 "color": s.get("color", "YELLOW"), "confidence": s.get("confidence", ""),
                 "proof": s.get("source_sentence", ""), "seg": sid}
        break

    drugs = []
    for d in ents.get("drugs", []):
        if d.get("negated"):
            continue
        dose = d.get("dose")
        dose = "NIL" if dose is None or dose == "" else (str(int(float(dose))) if str(dose).replace(".", "", 1).isdigit() and float(dose).is_integer() else str(dose))
        unit = d.get("unit", "") or "NIL"
        _, span = _seg_of(d.get("source_sentence"), segs)
        drugs.append({"name": d.get("name", "") or "NIL", "dose": dose,
                      "unit": unit, "frequency": d.get("frequency", "") or "NIL",
                      "duration": d.get("duration", "") or "NIL", "confidence": d.get("confidence", ""),
                      "color": d.get("color", "YELLOW"), "proof": f"seg {span} {d.get('source_sentence','')[:60]}",
                      "note": d.get("note", "")})

    vitals = ""
    if ents.get("vitals"):
        vitals = "; ".join(v.get("text", "") for v in ents["vitals"])

    # unmapped heard speech: BP-like text in transcript but no vitals extracted
    unmapped = ""
    if not vitals:
        import re
        m = re.search(r"BPA?[^.,]{0,30}", transcript_json.get("text", ""), re.I)
        if m:
            unmapped = m.group(0).strip()

    allergies_active = "; ".join(a.get("text", "") for a in ents.get("allergies", []) if not a.get("negated"))
    allergies_denied = [{"text": a.get("text", ""), "proof": a.get("source_sentence", "")}
                        for a in ents.get("allergies", []) if a.get("negated")]
    symptoms_denied = [{"text": s.get("text", ""), "proof": s.get("source_sentence", "")}
                       for s in ents.get("symptoms", []) if s.get("negated")]

    ctx = {"job_id": ents.get("job_id", transcript_json.get("job_id", "")),
           "language": transcript_json.get("language", ""),
           "template_id": spec["id"], "template_version": spec["version"],
           "date": datetime.date.today().isoformat(),
           "chief_complaint": chief, "hpi": transcript_json.get("text", "")[:500],
           "drugs": drugs, "vitals": vitals, "unmapped": unmapped,
           "allergies_active": allergies_active, "allergies_denied": allergies_denied,
           "symptoms_denied": symptoms_denied}
    html = tpl.render(**ctx)

    # DOCX fallback mirror (minimal styling)
    docx_bytes = None
    try:
        from docx import Document
        doc = Document()
        doc.add_heading(f"ER Discharge (DEMO) - {ctx['job_id']}", 1)
        doc.add_paragraph(f"lang: {ctx['language']} | template: {ctx['template_id']}:{ctx['template_version']} | {ctx['date']}")
        doc.add_heading("Chief Complaint", 2)
        doc.add_paragraph(chief.get("text", "___"))
        doc.add_heading("History", 2)
        doc.add_paragraph(ctx["hpi"])
        doc.add_heading("Drugs", 2)
        if drugs:
            t = doc.add_table(rows=1, cols=5)
            t.style = "Table Grid"
            for i, h in enumerate(["Name", "Dose", "Freq", "Days", "Conf"]):
                t.rows[0].cells[i].text = h
            for d in drugs:
                row = t.add_row().cells
                row[0].text = str(d["name"])
                row[1].text = f"{d['dose']} {d['unit']}"
                row[2].text = str(d["frequency"])
                row[3].text = str(d["duration"])
                row[4].text = f"{d['color']} {d['confidence']}"
        else:
            doc.add_paragraph("NIL — no drug heard")
        doc.add_heading("Vitals", 2)
        doc.add_paragraph(vitals or "NIL")
        doc.add_heading("Allergies (active)", 2)
        doc.add_paragraph(allergies_active or "None extracted.")
        for a in allergies_denied:
            doc.add_paragraph(f"DENIED: {a['text']} - {a['proof']}")
        doc.add_heading("Diagnosis (ICD-10)", 2)
        doc.add_paragraph("NIL — [PHYSICIAN TO CONFIRM]")
        doc.add_paragraph("Demo - not for clinical use.")
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()
    except Exception:
        docx_bytes = None

    return {"html": html, "docx_bytes": docx_bytes, "preview_data": ctx}


SOURCE_LIMIT = 2500  # chars of verbatim source text printed on the note (clearly marked if shortened)


def _review_counts(slots):
    from coords import AI_NOTE
    clin = [s for s in slots if not s["key"].startswith("hpi_inline") and s["key"] not in ("date_box", "sign_line")]
    return {"red": sum(1 for s in clin if s.get("color") == "RED"),
            "yellow": sum(1 for s in clin if s.get("color") == "YELLOW"),
            "ai": sum(1 for s in clin if AI_NOTE in str(s.get("note", "")))}


def _fill_premium(spec, entities_json, transcript_json, template_id, doc_status=None):
    """Discharge-note render: flowing A4 document (HTML, print-ready) + structured DOCX.

    Every medicine gets a table row, no clinical value is truncated, notes
    (conflicts, AI-suggested, multiple readings) stay visible, and the note
    states whether it is an AI-assisted draft or clinician-verified.
    """
    from jinja2 import Environment, FileSystemLoader
    from markupsafe import escape
    from coords import resolve_slots

    tpl_dir = os.path.join(_HERE, "templates")
    env = Environment(loader=FileSystemLoader(tpl_dir), autoescape=True)  # source text may contain < & "
    tpl = env.get_template(spec["layout"])
    r = resolve_slots(entities_json, transcript_json, template_id)
    tj = transcript_json or {}
    src = str(tj.get("text", "") or "")
    if len(src) > SOURCE_LIMIT:
        src = src[:SOURCE_LIMIT] + " … [source text shortened on the note — the complete text is in the MediBytes app]"
    status = {"state": "draft"}
    status.update(doc_status or {})
    ctx = {"job_id": (entities_json or {}).get("job_id", tj.get("job_id", "")),
           "language": tj.get("language", ""),
           "template_id": spec["id"], "template_version": spec["version"],
           "date": datetime.date.today().isoformat(),
           "slots": r["slots"], "S": {s["key"]: s for s in r["slots"]},
           "denied": r["denied"], "proof_lines": r["proof_lines"], "hpi_bg": r["hpi_bg"],
           "drug_rows": r["drug_rows"], "n_drugs": sum(1 for x in r["drug_rows"] if x["index"]),
           "review_notes": r["review_notes"], "denied_allergies": r["denied_allergies"],
           "denied_symptoms": r["denied_symptoms"], "symptoms_other": r["symptoms_other"],
           "unrecognised_meds": r["unrecognised_meds"], "other_vitals": r["other_vitals"],
           "source_text": src, "source_label": "image OCR" if tj.get("source") == "image" else "audio transcript",
           "doc_status": status, "review": _review_counts(r["slots"])}
    html = tpl.render(**ctx)
    # the rendered note must actually show every resolved value (allergies,
    # drugs, vitals, diagnosis...); a layout that drops them is not a note
    lost = [s["key"] for s in r["slots"]
            if not s["key"].startswith("hpi_inline") and str(escape(str(s.get("text", "")))) not in html]
    if not r["slots"] or lost:
        raise TemplateRenderError("discharge note NOT generated: layout "
                                  f"{spec['layout']} is missing clinical fields {lost or 'all'}")
    docx_error = None
    try:
        docx_bytes = _docx_word_editable(ctx)
    except Exception as e:  # HTML is complete; report the DOCX failure, never fake it
        docx_bytes, docx_error = None, f"{type(e).__name__}: {e}"
    return {"html": html, "docx_bytes": docx_bytes, "preview_data": ctx, "docx_error": docx_error}


def _docx_word_editable(ctx):
    """DOCX mirror of the HTML note: same sections, full medication table, visible notes.

    Value colours: RED = bold red (missing/conflict), YELLOW = highlighted
    (review / AI-suggested), GREEN = plain. Returns None if python-docx is absent.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_COLOR_INDEX
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
    except Exception:
        return None
    grey, red, navy = RGBColor(0x5B, 0x67, 0x70), RGBColor(0x99, 0x1B, 0x1B), RGBColor(0x0B, 0x25, 0x45)
    S = ctx["S"]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(15)
    sec.top_margin, sec.bottom_margin = Mm(14), Mm(15)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    def add_val(par, s, fallback="—"):
        if not s:
            r = par.add_run(fallback)
            r.font.color.rgb = grey
            return
        r = par.add_run(str(s.get("text", "")))
        if s.get("color") == "RED":
            r.font.bold, r.font.color.rgb = True, red
        elif s.get("color") == "YELLOW":
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if s.get("note"):
            n = par.add_run(f"  ({s['note']})")
            n.font.size, n.font.italic, n.font.color.rgb = Pt(8), True, grey

    def line(label, s=None, text=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        lab = p.add_run(f"{label}: ")
        lab.font.bold, lab.font.color.rgb = True, navy
        if text is not None:
            p.add_run(text)
        else:
            add_val(p, s)
        return p

    def shade(cell, hex_fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcpr.append(shd)

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(h)
            run.font.bold, run.font.size, run.font.color.rgb = True, Pt(9), RGBColor(0xFF, 0xFF, 0xFF)
            shade(c, "0B2545")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                if isinstance(v, dict) or v is None:
                    add_val(cells[i].paragraphs[0], v, fallback="")
                else:
                    cells[i].paragraphs[0].add_run(str(v))
        return t

    st = ctx["doc_status"]
    doc.add_heading(f"ER Discharge Summary — {ctx['job_id']}", 1)
    meta = doc.add_paragraph(f"MediBytes demonstration prototype · generated {ctx['date']} · source: "
                             f"{ctx['source_label']} · template {ctx['template_id']} v{ctx['template_version']}")
    meta.runs[0].font.size, meta.runs[0].font.color.rgb = Pt(8.5), grey
    sp = doc.add_paragraph()
    if st.get("state") == "verified":
        sr = sp.add_run(f"CLINICIAN-VERIFIED (DEMO) — verified {st.get('verified_at', '')} · "
                        f"{st.get('corrections', 0)} audited correction(s) in this run · note fingerprint "
                        f"{st.get('fingerprint', '')}. AI-suggested values keep their label. Not for clinical use.")
    else:
        rv = ctx["review"]
        sr = sp.add_run(f"AI-ASSISTED DRAFT — awaiting clinician verification. Not for clinical use. "
                        f"Review summary: {rv['red']} missing/conflicting (red) · {rv['yellow']} to review "
                        f"(yellow) · {rv['ai']} AI-suggested.")
    sr.font.bold = True
    lg = doc.add_paragraph("Colours: plain = read from the source by rules · highlighted = clinician review needed "
                           "(incl. AI-suggested) · bold red = missing or conflicting, resolve before verification.")
    lg.runs[0].font.size, lg.runs[0].font.color.rgb = Pt(8), grey

    doc.add_heading("Patient", 2)
    if S.get("patient_name"):
        line("Name", S["patient_name"])
    else:
        line("Name", text="____________________ (not captured)")
    line("Age / sex", text="____________________ (not captured)")

    doc.add_heading("Presenting complaint", 2)
    line("Chief complaint", S.get("cc_blank"))
    if ctx["symptoms_other"]:
        line("Other symptoms", text=", ".join(ctx["symptoms_other"]))
    if ctx["denied_symptoms"]:
        line("Denied symptoms", text=", ".join(ctx["denied_symptoms"]))

    doc.add_heading(f"Medications ({ctx['n_drugs']})", 2)
    table(["#", "Medicine", "Dose", "Unit", "Frequency", "Duration"],
          [[r["index"] or "—", r["cells"].get("name"), r["cells"].get("dose"), r["cells"].get("unit"),
            r["cells"].get("freq"), r["cells"].get("duration")] for r in ctx["drug_rows"]])
    if ctx["unrecognised_meds"]:
        w = doc.add_paragraph()
        wr = w.add_run("Possible medications not recognised — a dose was stated but the medicine name was not "
                       "recognised. Check the source and add manually if prescribed:")
        wr.font.bold, wr.font.color.rgb = True, red
        for u in ctx["unrecognised_meds"]:
            ur = doc.add_paragraph(style="List Bullet").add_run(f"“{u}”")
            ur.font.color.rgb = red

    doc.add_heading("Vital signs", 2)
    table(["Vital sign", "Value"], [["BP systolic (mmHg)", S.get("vitals_bp_sys")],
                                    ["BP diastolic (mmHg)", S.get("vitals_bp_dia")],
                                    ["Temperature", S.get("vitals_temp")], ["SpO2", S.get("vitals_spo2")]])
    if ctx["other_vitals"]:
        line("Other recorded vitals", text="; ".join(ctx["other_vitals"]))

    doc.add_heading("Allergies", 2)
    line("Allergy status", S.get("allergy_line"))
    if ctx["denied_allergies"]:
        line("Documented as not allergic", text=", ".join(ctx["denied_allergies"]))

    doc.add_heading("Diagnosis and follow-up", 2)
    line("Diagnosis", S.get("dx_text"))
    line("ICD-10 code", S.get("dx_code_box"))
    line("Follow-up", S.get("fu_text"))

    if ctx["review_notes"]:
        doc.add_heading("Items for clinician review", 2)
        for n in ctx["review_notes"]:
            doc.add_paragraph(n, style="List Bullet")

    doc.add_heading("Source " + ("OCR text" if ctx["source_label"] == "image OCR" else "transcript") + " (verbatim)", 2)
    src = doc.add_paragraph(ctx["source_text"])
    for r in src.runs:
        r.font.size, r.font.color.rgb = Pt(8.5), grey

    doc.add_heading("Clinician sign-off", 2)
    line("Signature", S.get("sign_line"))
    line("Note generated", S.get("date_box"))
    foot = doc.add_paragraph("Generated by MediBytes (demonstration prototype) from synthetic input. Every value must "
                             "be checked by a clinician before any use; red items must be completed or resolved.")
    foot.runs[0].font.size, foot.runs[0].font.color.rgb = Pt(8), grey
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

"""IPD P1 step 8: render a STORED progress-note DocumentVersion (HTML + DOCX).

A read-only view. It reads the selected DocumentVersion (the source of truth),
the Patient, Encounter, bed assignment and the cited events/captures for the
traceability table. It never regenerates, edits, approves or writes anything
and returns HTML (str) / DOCX (bytes) in memory.

- Line text is rendered exactly as stored (HTML autoescaped, not rewritten).
- Approval shown comes from the SELECTED version's approval fields only, not
  from Document.status, so an older approved version and a newer draft are
  never confused.
- Stale = stored stale flag OR a live documents.staleness() check (read-only).
- Missing patient/encounter data is shown as "Not available", never guessed.
"""
import io
import os
from datetime import datetime
from typing import Optional

from jinja2 import Environment, FileSystemLoader, StrictUndefined

from .documents import DocumentError, staleness
from .models import IST, DocumentType
from .store import IpdStore

try:
    import docx  # python-docx, already a demo dependency
    from docx.enum.text import WD_COLOR_INDEX
    DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover - environment dependent
    DOCX_AVAILABLE = False

TEMPLATE = "ipd_progress_note.html.j2"
_TEMPLATES = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates")
_ENV = Environment(loader=FileSystemLoader(_TEMPLATES), autoescape=True, undefined=StrictUndefined,
                   keep_trailing_newline=True)
DRAFT_MARKER = "DRAFT — NOT APPROVED FOR CLINICAL USE"
APPROVED_MARKER = "APPROVED"
STALE_MARKER = "STALE — REGENERATION REQUIRED"
REVIEW_MARKER = "CLINICIAN REVIEW REQUIRED"
NA = ""  # template shows "Not available"


def _fmt(ts: Optional[datetime]) -> str:
    return ts.astimezone(IST).strftime("%Y-%m-%d %H:%M IST") if ts else NA


def _bed_at(beds, at: datetime):
    """Bed assignment whose interval contains `at`; else the latest one before it; else None."""
    best = None
    for b in beds:
        if b.from_at <= at and (b.to_at is None or at < b.to_at):
            return b
        if b.from_at <= at:
            best = b
    return best


def build_render_context(store: IpdStore, document_id: str, version: Optional[int] = None) -> dict:
    """Plain dict for the template/DOCX. Read-only."""
    d = store.get_document(document_id)
    if d is None:
        raise DocumentError(f"unknown document {document_id}")
    if d.type != DocumentType.PROGRESS_NOTE:
        raise DocumentError(f"document {document_id} is not a progress note")
    v = store.get_document_version(d.id, version or d.current_version)
    if v is None:
        raise DocumentError(f"unknown version {version}")
    enc = store.get_encounter(d.encounter_id)
    pat = store.get_patient(enc.patient_id) if enc else None
    bed = _bed_at(store.list_bed_assignments(enc.id), v.created_at) if enc else None
    rep = staleness(store, d.id, v.version)
    c = v.content
    sections = [{"key": s["key"], "title": s["title"], "required": bool(s.get("required")),
                 "lines": [{"text": l["text"], "color": l["color"], "kind": l["kind"],
                            "needs_review": bool(l.get("needs_review")), "review_reason": l.get("review_reason", "")}
                           for l in s["lines"]]} for s in c.get("sections", [])]
    needs_review = bool(c.get("needs_review") or c.get("conflicts")) or any(
        l["needs_review"] for s in sections for l in s["lines"])

    trace = []
    for s in c.get("sections", []):
        for l in s["lines"]:
            if l.get("kind") == "placeholder":
                trace.append({"section": s["key"], "text": l["text"], "placeholder": True, "sources": []})
                continue
            sources = []
            for eid in l.get("event_ids", []):
                e = store.get_event(eid)
                if e is None:
                    sources.append({"event_id": eid, "source_type": "event not found", "capture": "event not found"})
                    continue
                cap = store.get_capture(e.source_capture_id) if e.source_capture_id else None
                if cap is not None:
                    capture = (f"{cap.source.value} / {cap.capture_context.value} by {cap.author_id} "
                               f"({cap.author_role.value}) at {_fmt(cap.captured_at)}")
                elif e.source_capture_id:
                    capture = f"capture {e.source_capture_id} not found"
                else:
                    capture = f"no capture (entered by {e.author_id}, {e.author_role.value})"
                sources.append({"event_id": eid, "source_type": e.source_type.value, "capture": capture})
            trace.append({"section": s["key"], "text": l["text"], "placeholder": False, "sources": sources})

    approved = v.approved_at is not None
    return {
        "patient": [("Name", pat.name if pat else NA), ("MRN", pat.mrn if pat else NA),
                    ("Date of birth", pat.dob.isoformat() if pat and pat.dob else NA),
                    ("Sex", pat.sex if pat and pat.sex else NA)],
        "encounter": [("Encounter ID", enc.id if enc else NA), ("Encounter type", enc.type.value if enc else NA),
                      ("Admitted", _fmt(enc.admit_at) if enc else NA),
                      ("Ward", bed.ward if bed else NA), ("Bed", bed.bed if bed else NA),
                      ("Unit", bed.unit_type.value if bed else NA)],
        "document": [("Document", "Progress Note"), ("Version", str(v.version)),
                     ("Status", "Approved" if approved else "Draft"),
                     ("Change type", v.change_type.value), ("Window from", _fmt(d.window_from)),
                     ("Window to", _fmt(d.window_to)), ("Created", _fmt(v.created_at)), ("Created by", v.created_by)]
                    + ([("Approved by", v.approved_by), ("Approver role", v.approved_role.value),
                        ("Approved at", _fmt(v.approved_at))] if approved else []),
        "doc": {"version": v.version, "approved": approved, "approved_by": v.approved_by or "",
                "approved_role": v.approved_role.value if v.approved_role else "", "approved_at": _fmt(v.approved_at),
                "stale": bool(v.stale or rep.stale), "stale_reasons": list(rep.reasons),
                "needs_review": needs_review, "change_type": v.change_type.value, "generator": v.generator},
        "sections": sections,
        "trace": trace,
    }


def render_progress_note_html(store: IpdStore, document_id: str, version: Optional[int] = None) -> str:
    return _ENV.get_template(TEMPLATE).render(**build_render_context(store, document_id, version))


def render_progress_note_docx(store: IpdStore, document_id: str, version: Optional[int] = None) -> bytes:
    """DOCX bytes with the same content and safety markers as the HTML."""
    if not DOCX_AVAILABLE:
        raise RuntimeError("python-docx is not installed; DOCX export unavailable")
    ctx = build_render_context(store, document_id, version)
    doc = docx.Document()
    props = doc.core_properties
    props.title, props.author, props.comments = "MediBytes IPD Progress Note", "MediBytes", ""
    doc.add_heading("MediBytes IPD Progress Note", level=0)
    d = ctx["doc"]
    banners = [f"{APPROVED_MARKER} — by {d['approved_by']} ({d['approved_role']}) at {d['approved_at']}"
               if d["approved"] else DRAFT_MARKER]
    if d["stale"]:
        banners.append(STALE_MARKER + (": " + "; ".join(d["stale_reasons"]) if d["stale_reasons"] else ""))
    if d["needs_review"]:
        banners.append(REVIEW_MARKER + " — conflicting values are shown below and have not been resolved")
    for b in banners:
        run = doc.add_paragraph().add_run(b)
        run.bold = True
    for title, rows in (("Patient", ctx["patient"]), ("Encounter", ctx["encounter"]), ("Document", ctx["document"])):
        doc.add_heading(title, level=1)
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for label, value in rows:
            cells = t.add_row().cells
            cells[0].text, cells[1].text = label, value or "Not available"
    doc.add_heading("SOAP Note", level=1)
    doc.add_paragraph("GREEN = verified · YELLOW = unverified, verify against source · RED = required information "
                      "missing · NIL = nothing captured (optional section)")
    hl = {"GREEN": WD_COLOR_INDEX.BRIGHT_GREEN, "YELLOW": WD_COLOR_INDEX.YELLOW, "RED": WD_COLOR_INDEX.RED,
          "NIL": WD_COLOR_INDEX.GRAY_25}
    for s in ctx["sections"]:
        doc.add_heading(f"{s['key']} — {s['title']}" + (" (required)" if s["required"] else ""), level=2)
        for l in s["lines"]:
            p = doc.add_paragraph()
            badge = p.add_run(f"[{l['color']}] ")
            badge.bold = True
            badge.font.highlight_color = hl.get(l["color"])
            txt = p.add_run(l["text"])
            txt.italic = l["kind"] == "placeholder"
            if l["needs_review"]:
                flag = p.add_run(" ⚠ REVIEW REQUIRED" + (f" — {l['review_reason']}" if l["review_reason"] else ""))
                flag.bold = True
    doc.add_heading("Clinical Traceability / Source Events", level=1)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for cell, h in zip(t.rows[0].cells, ("Section", "Line", "Source event(s)", "Source type", "Capture")):
        cell.text = h
    for tr in ctx["trace"]:
        cells = t.add_row().cells
        cells[0].text, cells[1].text = tr["section"], tr["text"]
        if tr["placeholder"]:
            cells[2].text, cells[3].text, cells[4].text = "Placeholder — no clinical source event", "", ""
        else:
            cells[2].text = "\n".join(s["event_id"] for s in tr["sources"])
            cells[3].text = "\n".join(s["source_type"] for s in tr["sources"])
            cells[4].text = "\n".join(s["capture"] for s in tr["sources"])
    doc.add_paragraph(f"Rendered from stored document version {d['version']} ({d['change_type']}). "
                      f"Generator: {d['generator']}. This document is a view of the stored version; "
                      "it does not change clinical records.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

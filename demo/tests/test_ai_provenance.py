"""Regression: AI-filled values must never look like confirmed (GREEN) facts.

Run:  python -m pytest demo/tests/test_ai_provenance.py -q

Before: AI gap-fills were written into the regex drug row, whose colour
(GREEN = dose+unit read from source) then coloured the AI value too, with no
per-field marker. Now every AI-provided value carries provenance
(ai_suggested / ai_fields / structured_vitals_ai), renders YELLOW with an
"AI-suggested - verify" note in the note, HTML and DOCX, and only a clinician
edit of that field removes the marker. Synthetic text only; LLM output faked.
"""
import io
import os
import sys

_HERE = os.path.dirname(os.path.abspath(__file__))
_DEMO = os.path.dirname(_HERE)
sys.path.insert(0, _DEMO)

from coords import resolve_slots  # noqa: E402
from fill_template import fill_template  # noqa: E402
from llm_extract import merge_primary  # noqa: E402
from ocr_llm_judge import merge_image  # noqa: E402
from stt_extract import _merge_tidy_drugs, run_text_extract  # noqa: E402

AI = "ai-suggested"
# "Twice daily after food" is not a frequency-only line, so regex leaves the
# frequency empty (a bare "Twice daily" line is now attached by regex) and the
# AI gap-fill under test is still needed.
OCR = ("Name: Test Person\nTab Paracetamol 500 mg\nTwice daily after food\nBP 120/80\n"
       "Saturation 99\nImpression: viral fever")
AUDIO = "Test cough. Give paracetamol 500 mg. Take it twice a day. BP 120/80."


def _run(text, source="audio"):
    segs = [{"id": i, "text": t, "start": 0.0, "end": 0.0, "lang": "", "confidence": 0.9,
             "words": []} for i, t in enumerate(text.split("\n"))]
    return run_text_extract(text, segs, "t", use_llm=False, source=source)


def _slots(ej, tj):
    return {s["key"]: s for s in resolve_slots(ej, tj)["slots"]}


def _docx_med_cell(doc, value):
    """Text of the DOCX medication-table cell holding `value` (structured DOCX layout)."""
    for t in doc.tables:
        if t.rows and t.rows[0].cells[1].text == "Medicine":
            for row in t.rows[1:]:
                for c in row.cells:
                    if value in c.text:
                        return c.text
    raise AssertionError(f"{value!r} not found in the DOCX medication table")


def _is_ai(slot):
    return AI in slot["note"].lower()


def _image_merged():
    tj, base = _run(OCR, "image")
    llm = {"drugs": [{"name": "paracetamol", "dose": 500, "unit": "mg", "frequency": "twice daily"}],
           "structured_vitals": {"spo2": "99"}, "vitals": [{"text": "99"}],
           "diagnosis": {"text": "viral fever", "confidence": 0.85, "color": "YELLOW"},
           "allergies": [{"text": "penicillin", "negated": False, "color": "YELLOW"}]}
    return tj, base, merge_image(base, llm, OCR + "\npenicillin rash")


# ---- reproduce: AI frequency was GREEN ----

def test_image_ai_frequency_is_review_not_green():
    tj, base, merged = _image_merged()
    s = _slots(merged, tj)
    assert s["drug_0_freq"]["text"] == "twice daily"  # value unchanged
    assert s["drug_0_freq"]["color"] == "YELLOW" and _is_ai(s["drug_0_freq"])
    # source-extracted cells of the same row keep their status
    assert s["drug_0_name"]["color"] == "GREEN" and not _is_ai(s["drug_0_dose"])
    assert s["drug_0_dose"]["color"] == "GREEN"


def test_audio_ai_frequency_is_review_not_green():
    tj, base = _run(AUDIO)
    assert base["drugs"][0]["frequency"] == ""  # regex did not read it
    merged = merge_primary(base, {"drugs": [{"name": "paracetamol", "frequency": "twice daily"}]})
    s = _slots(merged, tj)
    assert s["drug_0_freq"]["color"] == "YELLOW" and _is_ai(s["drug_0_freq"])
    assert s["drug_0_dose"]["color"] == "GREEN" and not _is_ai(s["drug_0_dose"])


def test_llm_tidy_fill_is_review_not_green():
    tj, base = _run(AUDIO)
    drugs, _ = _merge_tidy_drugs(base["drugs"], [{"name": "paracetamol", "frequency": "twice daily"}])
    s = _slots({**base, "drugs": drugs}, tj)
    assert s["drug_0_freq"]["color"] == "YELLOW" and _is_ai(s["drug_0_freq"])


def test_new_ai_drug_row_is_marked():
    tj, base = _run(AUDIO)
    merged = merge_primary(base, {"drugs": [{"name": "cetirizine", "dose": 10.0, "unit": "mg",
                                             "frequency": "", "duration": "", "color": "YELLOW",
                                             "note": "llm-primary"}]})
    s = _slots(merged, tj)
    assert s["drug_1_name"]["text"] == "cetirizine" and _is_ai(s["drug_1_name"])
    assert s["drug_1_dose"]["color"] != "GREEN"


# ---- other AI interpretations ----

def test_ai_vital_diagnosis_allergy_marked():
    tj, base, merged = _image_merged()
    s = _slots(merged, tj)
    assert s["vitals_spo2"]["text"] == "99" and _is_ai(s["vitals_spo2"])
    assert not _is_ai(s["vitals_bp_sys"])  # regex vital
    assert _is_ai(s["dx_text"]) and s["dx_text"]["color"] == "YELLOW"
    assert "penicillin" in s["allergy_line"]["text"] and _is_ai(s["allergy_line"])


def test_ai_patient_and_chief_complaint_marked():
    ocr = "Tab Paracetamol 500 mg twice daily\nThe patient Test Person reports fever"
    tj, base = _run(ocr, "image")
    merged = merge_image(base, {"drugs": [], "patient": {"name": "Test Person", "color": "YELLOW"},
                                "symptoms": [{"text": "headache", "negated": False}]},
                         ocr + "\nheadache")
    s = _slots(merged, tj)
    assert _is_ai(s["patient_name"])
    assert s["cc_blank"]["text"] == "fever" and not _is_ai(s["cc_blank"])  # regex symptom first


def test_ai_denied_symptoms_not_green():
    from llm_extract import extract_llm_primary
    import json
    import types
    import llm_extract
    orig = (llm_extract.ollama_available, llm_extract.subprocess.run)
    llm_extract.ollama_available = lambda model="x": (True, "")
    llm_extract.subprocess.run = lambda *a, **k: types.SimpleNamespace(
        stdout=json.dumps({"chief_complaint": "", "drugs": [], "vitals": {},
                           "symptoms_denied": ["chest pain"]}), stderr="")
    try:
        llm = extract_llm_primary("denies chest pain", "denies chest pain", [])
    finally:
        llm_extract.ollama_available, llm_extract.subprocess.run = orig
    assert all(x["color"] != "GREEN" for x in llm["symptoms"])


# ---- unchanged: regex-only, missing, conflicts ----

def test_regex_only_values_unchanged():
    tj, ej = _run("Give paracetamol 500 mg twice daily for 3 days. BP 120/80.")
    s = _slots(ej, tj)
    for k in ("drug_0_name", "drug_0_dose", "drug_0_freq", "drug_0_duration"):
        assert s[k]["color"] == "GREEN" and not _is_ai(s[k])
    assert s["drug_0_freq"]["text"] == "twice daily"


def test_missing_and_conflict_states_preserved():
    tj, base = _run(AUDIO)
    merged = merge_primary(base, {"drugs": [], "structured_vitals": {"sys": "140", "dia": "90"}})
    s = _slots(merged, tj)
    assert s["drug_0_duration"]["color"] == "RED"  # still missing
    assert "AI read 140" in s["vitals_bp_sys"]["note"]  # conflict kept, regex value kept
    assert s["vitals_bp_sys"]["text"] == "120"


# ---- clinician review ----

def test_clinician_edit_clears_marker_untouched_ai_stays():
    import app
    tj, base, merged = _image_merged()
    defaults = {k: v["text"] for k, v in _slots(merged, tj).items()}
    patched, audit = app._apply_edits(merged, {"drug_0_freq": "three times daily"}, defaults)
    s = _slots(patched, tj)
    assert s["drug_0_freq"]["text"] == "three times daily" and not _is_ai(s["drug_0_freq"])
    assert _is_ai(s["vitals_spo2"]) and _is_ai(s["dx_text"])  # not reviewed yet
    assert [a["slot_key"] for a in audit] == ["drug_0_freq"]


def test_clinician_edit_of_ai_vital_and_dx_clears_marker():
    import app
    tj, base, merged = _image_merged()
    defaults = {k: v["text"] for k, v in _slots(merged, tj).items()}
    patched, _ = app._apply_edits(merged, {"vitals_spo2": "98%", "dx_text": "viral fever (confirmed)"},
                                  defaults)
    s = _slots(patched, tj)
    assert not _is_ai(s["vitals_spo2"]) and not _is_ai(s["dx_text"])


def test_unrelated_edit_keeps_ai_marker():
    import app
    tj, base, merged = _image_merged()
    defaults = {k: v["text"] for k, v in _slots(merged, tj).items()}
    patched, _ = app._apply_edits(merged, {"fu_text": "Review in 3 days"}, defaults)
    assert _is_ai(_slots(patched, tj)["drug_0_freq"])


# ---- exports ----

def test_ai_marker_in_html_and_docx():
    from docx import Document
    tj, base, merged = _image_merged()
    out = fill_template(merged, tj)
    html = out["html"]
    i = html.index("twice daily")
    cell = html[html.rfind("<div", 0, i):html.find("</div>", i)]
    assert "gYELLOW" in cell and "AI-suggested" in cell
    cell = _docx_med_cell(Document(io.BytesIO(out["docx_bytes"])), "twice daily")
    assert "AI-suggested" in cell


# ---- app: provenance survives reruns, corrections, Verify and verified exports ----

def _app_script():
    import os
    import shutil
    import sys

    sys.path.insert(0, os.environ["MB_TEST_DEMO_DIR"])
    import streamlit as st
    import app

    class _Up:
        name = "synthetic_ai_fill.wav"

        def getvalue(self):
            return b"RIFF-synthetic-ai-fill"

        def getbuffer(self):
            return memoryview(self.getvalue())

    def fake_clean(i, o, apply_denoise=True):
        shutil.copyfile(i, o)
        return {"duration_s": 1.0, "vad_ratio": 1.0, "denoise": "none"}

    def fake_stt(clean_wav, job_id="t", use_llm="auto", model="", ollama_model=""):
        from llm_extract import merge_primary
        from stt_extract import run_text_extract
        text = "Test cough. Give paracetamol 500 mg. Take it twice a day. BP 120/80."
        tj, base = run_text_extract(text, [{"id": 0, "text": text, "start": 0.0, "end": 2.0,
                                            "lang": "en-IN", "confidence": 0.9, "words": []}],
                                    job_id, use_llm=False, engine="fake-test")
        ej = merge_primary(base, {"drugs": [{"name": "paracetamol", "frequency": "twice daily"}]})
        return tj, {"job_id": job_id, **ej}

    saved = (st.file_uploader, app.clean_audio, app.run_stt_extract, app._ollama_available, app._HERE)
    st.file_uploader = lambda *a, **k: _Up()
    app.clean_audio = fake_clean
    app.run_stt_extract = fake_stt
    app._ollama_available = lambda model: (False, "disabled in test")
    app._HERE = os.environ["MB_TEST_OUT_DIR"]
    try:
        app.main()
    finally:
        (st.file_uploader, app.clean_audio, app.run_stt_extract, app._ollama_available,
         app._HERE) = saved


def test_app_ai_marker_survives_rerun_verify_and_export(tmp_path, monkeypatch):
    from docx import Document
    from streamlit.testing.v1 import AppTest
    out = tmp_path / "out"
    for d in ("entities", "exports", "transcripts"):
        (out / d).mkdir(parents=True)
    monkeypatch.setenv("MB_TEST_DEMO_DIR", _DEMO)
    monkeypatch.setenv("MB_TEST_OUT_DIR", str(out))
    at = AppTest.from_function(_app_script, default_timeout=120)
    at.run()
    next(c for c in at.checkbox if c.label.startswith("LLM")).set_value(False)
    next(b for b in at.button if b.label == "▶ Run pipeline").click().run()
    assert not at.exception, at.exception
    job = next(str(s.value).split("`")[1] for s in at.success if str(s.value).startswith("Job `"))
    cards = " ".join(str(m.value) for m in at.markdown)
    assert "AI-suggested - verify:</b> frequency" in cards
    # clinician fills the missing fields but does NOT touch the AI frequency
    for sk, v in {"drug_0_duration": "3 days", "vitals_temp": "98.6 F", "vitals_spo2": "98%",
                  "dx_text": "test condition", "dx_code_box": "Z00.0"}.items():
        (at.text_area if sk == "dx_text" else at.text_input)(key=f"{job}_{sk}").set_value(v)
    next(b for b in at.button if b.label == "Apply corrections").click().run()
    at.run()  # harmless rerun
    next(b for b in at.button if b.key == f"verify_{job}").click().run()
    assert any("Verified" in str(s.value) for s in at.success)
    html = (out / "exports" / f"{job}.corrected.html").read_text(encoding="utf-8")
    i = html.index("twice daily")
    cell = html[html.rfind("<div", 0, i):html.find("</div>", i)]
    assert "gYELLOW" in cell and "AI-suggested" in cell  # provenance kept in the verified export
    assert "AI-suggested" in _docx_med_cell(Document(str(out / "exports" / f"{job}.corrected.docx")),
                                             "twice daily")

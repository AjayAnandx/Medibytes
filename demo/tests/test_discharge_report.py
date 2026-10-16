"""Regression: ER discharge-note output (HTML + DOCX) must be complete and honest.

Run:  python -m pytest demo/tests/test_discharge_report.py -q

Defects found on 2026-10-16 by rendering synthetic notes and printing them to A4:
only 2 medicines printed (3rd+ silently dropped), values cut with "…",
notes (AI-suggested / conflict / multiple readings) clipped, unlabelled vitals,
overlapping static text, no draft/verified status, unescaped source text,
missing follow-up shown GREEN, DOCX a developer key list without medicines 3+.
Synthetic text only; LLM output is a fixed fake dict (no Ollama).
"""
import io
import os
import re
import sys

import pytest

_HERE = os.path.dirname(os.path.abspath(__file__))
_DEMO = os.path.dirname(_HERE)
sys.path.insert(0, _DEMO)

from coords import resolve_slots  # noqa: E402
from fill_template import fill_template  # noqa: E402
from llm_extract import merge_primary  # noqa: E402
from stt_extract import _load_drug_list, run_text_extract  # noqa: E402

LONG_DX = ("community acquired pneumonia of the right lower lobe with reactive airway component, "
           "clinically stable for outpatient management")
LONG_FU = ("Follow up in the respiratory clinic after one week with a repeat chest X-ray and full blood count, "
           "or return to the emergency department immediately if breathing worsens")
STRESS = ("Patient has productive cough and fever for five days with mild breathlessness. "
          "BP 150/95 on arrival. Repeat BP 130/80 after rest. Temp 101 degrees Fahrenheit. SpO2 94% on room air. "
          "Give amoxicillin 875 mg twice daily for 7 days. Give paracetamol 500 mg every 6 hours for 5 days. "
          "Give salbutamol 100 mcg every 4 hours for 10 days. Give cetirizine 10 mg. "
          "Take the cetirizine once daily for 7 days. "
          "No known allergies. Patient is allergic to penicillin and sulfa. "
          f"Diagnosis: {LONG_DX}. {LONG_FU}.")


def _note(text, ai=False):
    segs = [{"id": 0, "text": text, "start": 0.0, "end": 30.0, "lang": "en-IN", "confidence": 0.9, "words": []}]
    tj, ej = run_text_extract(text, segs, "t-dn", use_llm=False)
    if ai:
        ej = {"job_id": "t-dn", **merge_primary(ej, {"drugs": [{"name": "cetirizine", "frequency": "once daily",
                                                                "duration": "7 days"}]})}
    return tj, ej


def _docx_text(b):
    from docx import Document
    d = Document(io.BytesIO(b))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _visible(html):
    """Text content of the HTML body (no tags, no CSS)."""
    body = html.split("<body>", 1)[-1]
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", body)).replace("&amp;", "&").replace("&lt;", "<")


# ---- every medicine, nothing cut ----

@pytest.mark.parametrize("n", [3, 4, 12, 28])
def test_every_medicine_printed_in_html_and_docx(n):
    names = [d["name"] for d in _load_drug_list() if d["name"] not in ("penicillin", "sulfa")][:n]
    text = "Cough. " + " ".join(f"Give {x} 100 mg twice daily for 5 days." for x in names)
    tj, ej = _note(text)
    assert [d["name"] for d in ej["drugs"]] == names
    out = fill_template(ej, tj)
    vis, docx = _visible(out["html"]), _docx_text(out["docx_bytes"])
    for x in names:  # was: only the first 2 medicines on the note and in the DOCX
        assert x in vis, x
        assert x in docx, x
    assert f"Medications ({n})" in vis and f"Medications ({n})" in docx


def test_long_values_are_never_truncated():
    tj, ej = _note(STRESS)
    out = fill_template(ej, tj)
    vis, docx = _visible(out["html"]), _docx_text(out["docx_bytes"])
    for full in (LONG_DX, LONG_FU, "101 degrees Fahrenheit", "penicillin; sulfa"):
        assert full in vis, full  # was "…right lower l…", "…with a r…", "101 degrees Fahren"
        assert full in docx, full
    assert "…" not in "".join(s["text"] for s in resolve_slots(ej, tj)["slots"])


def test_safety_notes_are_visible_in_both_formats():
    tj, ej = _note(STRESS, ai=True)
    out = fill_template(ej, tj)
    vis, docx = _visible(out["html"]), _docx_text(out["docx_bytes"])
    for needle in ("AI-suggested - verify", "multiple readings - verify: 150/95; 130/80",
                   "CONFLICT - clinician review"):
        assert needle in vis, needle  # was clipped by fixed-height boxes
        assert needle in docx, needle
    assert "Items for clinician review" in vis and "ALLERGY CONFLICT" in vis


def test_vitals_are_labelled():
    tj, ej = _note(STRESS)
    vis = _visible(fill_template(ej, tj)["html"])
    for label, value in (("BP systolic (mmHg)", "130"), ("BP diastolic (mmHg)", "80"),
                         ("Temperature", "101 degrees Fahrenheit"), ("SpO2", "94%")):
        i = vis.index(label)
        assert value in vis[i:i + 160], label


def test_other_and_denied_symptoms_listed():
    tj, ej = _note("Cough, fever and headache. Patient denies chest pain and breathlessness. "
                   "Give paracetamol 500 mg twice daily for 3 days.")
    vis = _visible(fill_template(ej, tj)["html"])
    assert "Other symptoms" in vis and "headache" in vis
    assert "Denied symptoms" in vis and "chest pain" in vis and "breathlessness" in vis


# ---- status, provenance and honesty ----

def test_draft_status_and_review_summary():
    tj, ej = _note(STRESS, ai=True)
    out = fill_template(ej, tj)
    vis, docx = _visible(out["html"]), _docx_text(out["docx_bytes"])
    assert "AI-ASSISTED DRAFT" in vis and "AI-ASSISTED DRAFT" in docx
    assert re.search(r"\b2 AI-suggested\b", vis) and re.search(r"\b2 AI-suggested\b", docx)
    assert not re.search(r"\bverified\b", out["html"], re.I)  # never claims verification


def test_verified_status_shows_who_when_and_corrections():
    tj, ej = _note(STRESS)
    st = {"state": "verified", "verified_at": "2026-10-16T10:15:00", "corrections": 3,
          "fingerprint": "abc123def4567890"}
    out = fill_template(ej, tj, doc_status=st)
    for text in (_visible(out["html"]), _docx_text(out["docx_bytes"])):
        assert "CLINICIAN-VERIFIED" in text and "2026-10-16T10:15:00" in text
        assert "3 audited correction(s)" in text and "abc123def4567890" in text
        assert "AI-ASSISTED DRAFT" not in text


def test_missing_values_stay_missing_and_followup_is_a_review_item():
    tj, ej = _note("Test cough.")
    slots = {s["key"]: s for s in resolve_slots(ej, tj)["slots"]}
    assert slots["drug_0_name"]["text"].startswith("NIL") and slots["drug_0_name"]["color"] == "RED"
    assert slots["fu_text"]["color"] == "YELLOW"  # was GREEN ("no follow-up heard")
    assert "no follow-up documented" in slots["fu_text"]["text"]
    vis = _visible(fill_template(ej, tj)["html"])
    assert "NOT DOCUMENTED" in vis and "No known allergies" not in vis


def test_no_real_organisation_name_in_header():
    tj, ej = _note(STRESS)
    out = fill_template(ej, tj)
    assert "CITY CARE HOSPITAL" not in out["html"].upper()
    assert "MediBytes demonstration prototype" in _visible(out["html"])


def test_source_text_is_escaped():
    text = "Dizziness. BP <90 & falling per triage note. Give paracetamol 500 mg twice daily for 3 days."
    tj, ej = _note(text)
    html = fill_template(ej, tj)["html"]
    assert "BP &lt;90 &amp; falling" in html and "BP <90 &" not in html


def test_long_source_text_is_marked_when_shortened():
    text = "Cough. Give paracetamol 500 mg twice daily for 3 days. " + "Supportive care discussed. " * 200
    tj, ej = _note(text)
    vis = _visible(fill_template(ej, tj)["html"])
    assert "source text shortened on the note" in vis  # never silently cut


def test_html_and_docx_carry_every_slot_value():
    tj, ej = _note(STRESS, ai=True)
    out = fill_template(ej, tj)
    vis, docx = _visible(out["html"]), _docx_text(out["docx_bytes"])
    for s in resolve_slots(ej, tj)["slots"]:
        if s["key"].startswith("hpi_inline"):
            continue
        assert s["text"] in vis, s["key"]
        assert s["text"] in docx, s["key"]


# ---- editor and AI judge cover every medicine row ----

def test_editor_exposes_and_applies_third_medicine():
    import app
    tj, ej = _note(STRESS)
    defaults = {s["key"]: s["text"] for s in resolve_slots(ej, tj)["slots"]}
    keys = app._editable_keys(defaults)
    assert "drug_2_name" in keys and "drug_3_duration" in keys
    assert keys.index("drug_1_duration") < keys.index("drug_2_name") < keys.index("vitals_bp_sys")
    assert app._label("drug_2_freq") == "Drug 3 frequency"
    patched, audit = app._apply_edits(ej, {"drug_2_dose": "200"}, defaults)
    assert [a["slot_key"] for a in audit] == ["drug_2_dose"]
    s = {x["key"]: x for x in resolve_slots(patched, tj)["slots"]}
    assert s["drug_2_dose"]["text"] == "200" and s["drug_2_name"]["text"] == "salbutamol"


def test_ai_judge_evaluates_every_medicine_row(monkeypatch):
    import json
    import types
    import ocr_llm_judge as J
    monkeypatch.setattr(J, "ollama_available", lambda model="x": (True, ""))
    monkeypatch.setattr("subprocess.run", lambda *a, **k: types.SimpleNamespace(
        stdout=json.dumps({"fields": []}), stderr=""))
    tj, ej = _note(STRESS)
    slots = resolve_slots(ej, tj)["slots"]
    ev = J.judge_fields(slots, STRESS, [{"line": 0, "text": STRESS}])
    fields = {f["field"] for f in ev["fields"]}
    assert {"drug_2_name", "drug_3_freq"} <= fields


# ---- through the real Streamlit editor: correct medicine 3, verify, export ----

from test_verification_state import _apply, _click, _job, _start, _verify, env  # noqa: E402,F401

FOUR = ("Productive cough for five days. BP 130/80. Temp 100.4 degrees Fahrenheit. SpO2 96% on room air. "
        "Give amoxicillin 875 mg twice daily for 7 days. Give paracetamol 500 mg every 6 hours for 5 days. "
        "Give salbutamol 100 mcg every 4 hours for 10 days. Give cetirizine 10 mg once daily for 7 days. "
        "Patient is allergic to sulfa. Diagnosis: acute bronchitis. Follow up in one week at the clinic.")


def test_app_edit_third_medicine_verify_and_export(env):
    configure, out = env
    configure(note=FOUR, icd="J20.9")
    at = _start()
    _click(at, "▶ Run pipeline")
    job = _job(at)
    assert at.text_input(key=f"{job}_drug_2_dose").value == "100"  # 3rd medicine is editable now
    _apply(at, job, drug_2_dose="200")
    _verify(at, job)
    html = (out / "exports" / f"{job}.corrected.html").read_text(encoding="utf-8")
    vis = _visible(html)
    assert "CLINICIAN-VERIFIED" in vis and "1 audited correction(s)" in vis
    for name in ("amoxicillin", "paracetamol", "salbutamol", "cetirizine"):
        assert name in vis
    i = vis.index("salbutamol")
    assert "200" in vis[i:i + 40]
    docx = _docx_text((out / "exports" / f"{job}.corrected.docx").read_bytes())
    assert "CLINICIAN-VERIFIED" in docx and "salbutamol" in docx and "cetirizine" in docx

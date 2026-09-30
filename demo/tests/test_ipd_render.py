"""IPD P1 step 8: read-only rendering of stored progress-note versions (HTML + DOCX).

Run:  python -m unittest demo/tests/test_ipd_render.py -v   (from repo root)
In-memory SQLite, synthetic data. Output stays in memory.
"""
import io
import os
import re
import sys
import unittest
from datetime import date, datetime, timedelta

_HERE = os.path.dirname(os.path.abspath(__file__))
_DEMO = os.path.dirname(_HERE)
sys.path.insert(0, _DEMO)

from ipd import render as render_mod  # noqa: E402
from ipd.capture import capture_typed_ward_round  # noqa: E402
from ipd.documents import approve, create_progress_note_document, regenerate_progress_note  # noqa: E402
from ipd.encounter import admit_patient  # noqa: E402
from ipd.models import (  # noqa: E402
    IST, DocumentStatus, Encounter, Event, EventCategory, Patient, Role, SourceType, Verification,
)
from ipd.progress_note import generate_progress_note  # noqa: E402
from ipd.render import (  # noqa: E402
    APPROVED_MARKER, DOCX_AVAILABLE, DRAFT_MARKER, REVIEW_MARKER, STALE_MARKER, build_render_context,
    render_progress_note_docx, render_progress_note_html,
)
from ipd.store import IpdStore  # noqa: E402
from ipd.timeline import Timeline  # noqa: E402

T0 = datetime(2026, 10, 6, 8, 0, tzinfo=IST)
H = timedelta(hours=1)
DR = ("u_dr_demo", Role.CONSULTANT)
NOTE = ("Patient has fever. No chest pain.\nBP 118/76, pulse 82.\nTemperature: 101 F\n"
        "Diagnosis: viral fever.\nStart paracetamol 650 mg TDS.")


def soap_part(html):
    return html.split("<h2>SOAP Note</h2>")[1].split("<h2>Clinical Traceability")[0]


def trace_part(html):
    return html.split("<h2>Clinical Traceability / Source Events</h2>")[1]


def docx_text(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


class RenderBase(unittest.TestCase):
    def setUp(self):
        self.store = IpdStore()
        self.tl = Timeline(self.store)
        pat = Patient(mrn="DEMO-R01", name="Demo Patient R01", dob=date(1980, 1, 15), sex="F")
        self.adm = admit_patient(self.store, pat, ward="Ward 4", bed="12", admitted_by=DR[0], admitted_role=DR[1],
                                 admit_at=T0)
        self.enc = self.adm.encounter.id
        self.cap = capture_typed_ward_round(self.store, encounter_id=self.enc, text=NOTE, author_id=DR[0],
                                            author_role=DR[1], location={"ward": "Ward 4", "bed": "12"},
                                            captured_at=T0 + 2 * H)
        self.doc, self.v1 = self.make_doc(self.enc)

    def tearDown(self):
        self.store.close()

    def make_doc(self, enc, at=T0 + 3 * H):
        d = generate_progress_note(self.tl, enc, window_from=T0, window_to=T0 + 12 * H, generated_by=DR[0],
                                   generated_at=at)
        return create_progress_note_document(self.store, d, created_by=DR[0], created_role=DR[1])

    def html(self, doc=None, version=None):
        return render_progress_note_html(self.store, (doc or self.doc).id, version)

    def lines(self, version):
        return [(s["key"], l) for s in version.content["sections"] for l in s["lines"]]


class TestHeaderAndSoap(RenderBase):
    def test_patient_encounter_bed_version(self):
        h = self.html()
        for text in ("MediBytes IPD Progress Note", "Demo Patient R01", "DEMO-R01", "1980-01-15", ">F<",
                     self.enc, "2026-10-06 08:00 IST", "Ward 4", ">12<", "WARD", "Progress Note",
                     "<th>Version</th><td>1</td>", "2026-10-06 11:00 IST"):
            self.assertIn(text, h)

    def test_soap_sections_render_with_exact_wording(self):
        h = soap_part(self.html())
        for key, title in (("S", "Subjective"), ("O", "Objective"), ("A", "Assessment"), ("P", "Plan")):
            self.assertIn(f"{key} — {title}", h)
        for _, l in self.lines(self.v1):
            self.assertIn(f">{l['text']}<", h)  # exactly as stored, nothing rewritten
        self.assertIn(">Fever present.<", h)
        self.assertIn(">10:00 Temperature 101 F<", h)
        self.assertIn(">Provisional diagnosis: viral fever.<", h)

    def test_proposed_medication_stays_proposed(self):
        h = self.html()
        self.assertIn("Proposed medication (not confirmed, not prescribed): paracetamol 650 mg TDS", h)
        for m in re.finditer(r"prescribed", h):
            self.assertEqual(h[m.start() - 4:m.start()], "not ")
        for bad in ("administered", "Active medication", "Rx:"):
            self.assertNotIn(bad, h)

    def test_event_ids_only_in_traceability(self):
        h = self.html()
        ids = self.v1.source_event_ids
        for eid in ids:
            self.assertNotIn(eid, soap_part(h))
            self.assertIn(eid, trace_part(h))
        self.assertIn("typed", trace_part(h))
        self.assertIn("text / ward_round by u_dr_demo (CONSULTANT)", trace_part(h))

    def test_escaping(self):
        self.assertNotIn("<script", self.html())
        e = Event(encounter_id=self.enc, occurred_at=T0 + 4 * H, category=EventCategory.PLAN, subtype="plan",
                  payload={"text": "<b>x</b> & y"}, author_id=DR[0], author_role=DR[1], source_type=SourceType.TYPED)
        self.tl.append(e)
        doc, _ = self.make_doc(self.enc, at=T0 + 5 * H)
        h = self.html(doc)
        self.assertIn("Plan: &lt;b&gt;x&lt;/b&gt; &amp; y.", h)
        self.assertNotIn("<b>x</b>", h)


class TestMarkers(RenderBase):
    def test_draft_marker(self):
        h = self.html()
        self.assertIn(DRAFT_MARKER, h)
        self.assertNotIn('class="banner approved"', h)

    def test_approved_marker_and_metadata(self):
        approve(self.store, self.doc.id, version=1, approved_by="u_cons_demo", approved_role=Role.RESIDENT,
                approved_at=T0 + 4 * H)
        h = self.html()
        self.assertIn(f"{APPROVED_MARKER} — by u_cons_demo (RESIDENT) at 2026-10-06 12:00 IST", h)
        for text in ("<th>Approved by</th><td>u_cons_demo</td>", "<th>Approver role</th><td>RESIDENT</td>",
                     "<th>Approved at</th><td>2026-10-06 12:00 IST</td>"):
            self.assertIn(text, h)
        self.assertNotIn(DRAFT_MARKER, h)

    def test_selected_version_approval_not_document_status(self):
        approve(self.store, self.doc.id, version=1, approved_by="u_cons_demo", approved_role=Role.CONSULTANT,
                approved_at=T0 + 4 * H)
        self.tl.append(Event(encounter_id=self.enc, occurred_at=T0 + 5 * H, category=EventCategory.PLAN,
                             subtype="followup", payload={"text": "Review tomorrow"}, author_id=DR[0],
                             author_role=DR[1], source_type=SourceType.TYPED))
        regenerate_progress_note(self.store, self.doc.id, generated_by=DR[0], generated_role=DR[1],
                                 generated_at=T0 + 6 * H)
        self.assertIn(APPROVED_MARKER + " — by", self.html(version=1))   # older approved version
        self.assertIn(DRAFT_MARKER, self.html(version=2))               # newer current draft
        self.assertIn(DRAFT_MARKER, self.html())                        # default = current
        # Document.status alone does not make an unapproved version look approved
        self.store.update_document(self.doc.id, 2, DocumentStatus.APPROVED)
        self.assertIn(DRAFT_MARKER, self.html(version=2))

    def test_colour_markers(self):
        tl = self.tl
        temp = next(e for e in tl.active_events(self.enc) if e.subtype == "temperature")
        tl.correct(temp.id, payload={"value": 100.4, "unit": "F"}, reason="re-measured", author_id=DR[0],
                   author_role=DR[1])
        doc, _ = self.make_doc(self.enc, at=T0 + 4 * H)
        h = soap_part(self.html(doc))
        self.assertIn('<li class="GREEN"><span class="badge GREEN">GREEN</span><span class="">10:00 Temperature 100.4 F', h)
        self.assertIn('<li class="YELLOW"><span class="badge YELLOW">YELLOW</span><span class="">Fever present.', h)
        self.assertIn('<li class="YELLOW"><span class="badge YELLOW">YELLOW</span><span class="">Proposed medication', h)
        self.assertNotIn('class="badge GREEN">GREEN</span><span class="">Fever', h)  # unverified stays YELLOW

    def test_red_and_nil_markers(self):
        other = admit_patient(self.store, Patient(mrn="DEMO-R02", name="Demo Patient R02"), ward="Ward 5", bed="1",
                              admitted_by=DR[0], admitted_role=DR[1], admit_at=T0).encounter.id
        doc, _ = self.make_doc(other)
        h = soap_part(self.html(doc))
        self.assertIn('<li class="RED"><span class="badge RED">RED</span><span class="placeholder">No subjective '
                      'information captured.', h)
        self.assertIn('<li class="RED"><span class="badge RED">RED</span><span class="placeholder">No assessment '
                      'information captured.', h)
        self.assertIn('<li class="NIL"><span class="badge NIL">NIL</span><span class="placeholder">No objective '
                      'findings captured.', h)
        self.assertIn('<li class="NIL"><span class="badge NIL">NIL</span><span class="placeholder">No plan captured.', h)

    def test_placeholder_has_no_provenance(self):
        other = admit_patient(self.store, Patient(mrn="DEMO-R03", name="Demo Patient R03"), ward="Ward 5", bed="2",
                              admitted_by=DR[0], admitted_role=DR[1], admit_at=T0).encounter.id
        doc, _ = self.make_doc(other)
        t = trace_part(self.html(doc))
        self.assertEqual(t.count("Placeholder — no clinical source event"), 4)
        self.assertNotIn("<code>", t)
        ctx = build_render_context(self.store, doc.id)
        self.assertTrue(all(tr["placeholder"] and tr["sources"] == [] for tr in ctx["trace"]))

    def test_conflict_review_warning(self):
        self.tl.append(Event(encounter_id=self.enc, occurred_at=T0 + 2 * H, category=EventCategory.VITAL,
                             subtype="pulse", payload={"value": 118}, author_id="dev_monitor", author_role=Role.DEVICE,
                             source_type=SourceType.DEVICE, confidence=1.0, verification=Verification.AUTO))
        doc, _ = self.make_doc(self.enc, at=T0 + 4 * H)
        h = self.html(doc)
        self.assertIn(REVIEW_MARKER, h)
        self.assertIn(">10:00 Pulse 82<", h)
        self.assertIn(">10:00 Pulse 118<", h)  # both values still visible
        self.assertEqual(soap_part(h).count("⚠ REVIEW REQUIRED"), 2)
        self.assertIn("device", trace_part(h))

    def test_stale_warning(self):
        self.assertNotIn(STALE_MARKER, self.html())
        fever = next(e for e in self.tl.active_events(self.enc) if e.subtype == "present")
        self.tl.retract(fever.id, reason="wrong patient", author_id=DR[0], author_role=DR[1])
        h = self.html()
        self.assertIn(STALE_MARKER, h)
        self.assertIn(f"event {fever.id} is retracted", h)
        self.assertIn(">Fever present.<", h)  # stored content still shown as stored, flagged stale


class TestSafety(RenderBase):
    def test_missing_patient_metadata(self):
        p = self.store.add_patient(Patient(mrn="DEMO-R04", name="Demo Patient R04"))  # no DOB / sex
        enc = self.store.add_encounter(Encounter(patient_id=p.id, admit_at=T0))      # no bed assignment
        doc, _ = self.make_doc(enc.id)
        h = self.html(doc)
        for label in ("Date of birth", "Sex", "Ward", "Bed", "Unit"):
            self.assertIn(f'<th>{label}</th><td><span class="na">Not available</span></td>', h)

    def test_renderer_does_not_mutate(self):
        tables = ("patients", "encounters", "bed_assignments", "captures", "source_records", "events", "documents",
                  "document_versions", "audit_log")
        snap = lambda: {t: [tuple(r) for r in self.store._conn.execute(f"SELECT * FROM {t}")] for t in tables}  # noqa
        before = snap()
        self.html()
        self.html(version=1)
        if DOCX_AVAILABLE:
            render_progress_note_docx(self.store, self.doc.id)
        build_render_context(self.store, self.doc.id)
        self.assertEqual(snap(), before)

    def test_deterministic(self):
        self.assertEqual(self.html(), self.html())
        self.assertEqual(build_render_context(self.store, self.doc.id), build_render_context(self.store, self.doc.id))

    def test_no_network_llm_or_scripts(self):
        import ast
        with open(render_mod.__file__, encoding="utf-8") as fh:
            tree = ast.parse(fh.read())
        mods = set()
        for n in ast.walk(tree):
            if isinstance(n, ast.Import):
                mods |= {a.name for a in n.names}
            elif isinstance(n, ast.ImportFrom):
                mods.add(("." * n.level) + (n.module or ""))
        self.assertEqual(mods, {"io", "os", "datetime", "typing", "jinja2", ".documents", ".models", ".store",
                                "docx", "docx.enum.text"})
        h = self.html()
        for bad in ("<script", "http://", "https://", " src=", "<link", "@import", "url("):
            self.assertNotIn(bad, h)

    def test_no_files_written(self):
        dirs = ("exports", "entities", "transcripts", "cleaned", "_state", "templates")

        def snap():
            s = {}
            for d in dirs:
                for base, _, files in os.walk(os.path.join(_DEMO, d)):
                    for f in files:
                        st = os.stat(os.path.join(base, f))
                        s[os.path.join(base, f)] = (st.st_size, st.st_mtime_ns)
            return s
        before = snap()
        self.html()
        if DOCX_AVAILABLE:
            render_progress_note_docx(self.store, self.doc.id)
        self.assertEqual(snap(), before)

    def test_unknown_document_or_version(self):
        from ipd.documents import DocumentError
        with self.assertRaises(DocumentError):
            render_progress_note_html(self.store, "doc_missing")
        with self.assertRaises(DocumentError):
            render_progress_note_html(self.store, self.doc.id, 99)


@unittest.skipUnless(DOCX_AVAILABLE, "python-docx not installed")
class TestDocx(RenderBase):
    def test_docx_content_and_markers(self):
        data = render_progress_note_docx(self.store, self.doc.id)
        self.assertTrue(data.startswith(b"PK"))
        text = docx_text(data)
        for s in ("MediBytes IPD Progress Note", DRAFT_MARKER, "Demo Patient R01", "DEMO-R01", self.enc, "Ward 4",
                  "S — Subjective", "O — Objective", "A — Assessment", "P — Plan", "Fever present.",
                  "Proposed medication (not confirmed, not prescribed): paracetamol 650 mg TDS",
                  "Clinical Traceability / Source Events"):
            self.assertIn(s, text)
        for eid in self.v1.source_event_ids:
            self.assertIn(eid, text)

    def test_docx_approved_and_review(self):
        approve(self.store, self.doc.id, version=1, approved_by="u_cons_demo", approved_role=Role.CONSULTANT,
                approved_at=T0 + 4 * H)
        text = docx_text(render_progress_note_docx(self.store, self.doc.id))
        self.assertIn("APPROVED — by u_cons_demo (CONSULTANT)", text)
        self.assertNotIn(DRAFT_MARKER, text)
        other = admit_patient(self.store, Patient(mrn="DEMO-R05", name="Demo Patient R05"), ward="Ward 5", bed="3",
                              admitted_by=DR[0], admitted_role=DR[1], admit_at=T0).encounter.id
        doc, _ = self.make_doc(other)
        t2 = docx_text(render_progress_note_docx(self.store, doc.id))
        self.assertIn("[RED] No subjective information captured.", t2)
        self.assertIn("[NIL] No plan captured.", t2)
        self.assertIn("Placeholder — no clinical source event", t2)


if __name__ == "__main__":
    unittest.main()
